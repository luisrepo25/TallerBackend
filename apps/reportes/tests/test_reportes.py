import io
import uuid
from datetime import date, timedelta
from decimal import Decimal

import openpyxl
from django.contrib.gis.geos import Polygon
from django.test import TestCase
from rest_framework import status
from rest_framework.test import APIClient

from apps.campanias.services import crear_campania
from apps.catastro.models import CatalogoTipoPredio
from apps.catastro.services import crear_predio
from apps.fiscalizacion.models import (
    ActaInspeccion,
    CatalogoEstadoExtintor,
    CatalogoInfraccion,
    CatalogoMaterialCombustible,
    Infraccion,
    MaterialRegistrado,
)
from apps.reportes.models import ReporteGenerado
from apps.usuarios.models import Rol
from apps.usuarios.services import crear_usuario


def cuadrado(x, y, lado=0.001):
    return Polygon(
        ((x, y), (x + lado, y), (x + lado, y + lado), (x, y + lado), (x, y)),
        srid=4326,
    )


class ReportesTestCase(TestCase):
    def setUp(self):
        self.suffix = uuid.uuid4().hex[:6]
        self.admin = self._crear_usuario("Administrador", "administrador", "admin")
        self.oficial = self._crear_usuario("Oficial de Mando", "oficial_mando", "oficial")
        self.inspector = self._crear_usuario("Inspector Tecnico", "inspector_tecnico", "insp")

        self.campania = crear_campania(
            nombre=f"Campaña Test {self.suffix}",
            fecha_inicio=date.today(),
            fecha_fin=date.today() + timedelta(days=30),
            creado_por=self.admin.administrador,
            geom=cuadrado(-63.18, -17.78, 0.05),
        )

        self.tipo_predio = CatalogoTipoPredio.objects.first()
        self.predio = crear_predio(
            nombre=f"Predio Comercial {self.suffix}",
            tipo_predio=self.tipo_predio,
            geom=cuadrado(-63.17, -17.77),
            area_m2=Decimal("250.00"),
            aforo=120,
            registrado_por=self.admin,
        )

        self.estado_extintor = CatalogoEstadoExtintor.objects.first()
        self.acta = ActaInspeccion.objects.create(
            predio=self.predio,
            campania=self.campania,
            inspector=self.inspector.inspectortecnico,
            estado_extintores=self.estado_extintor,
            rutas_evacuacion_despejadas=True,
            fallas_electricas=False,
            distancia_hidrante_m=Decimal("85.50"),
            carga_fuego=Decimal("120.40"),
            scoring_riesgo=Decimal("35.50"),
            cluster_riesgo="Bajo",
            cerrada=True,
        )

        # Registrar material e infraccion de prueba
        mat = CatalogoMaterialCombustible.objects.first()
        if mat:
            MaterialRegistrado.objects.create(
                acta=self.acta,
                material=mat,
                peso_kg=Decimal("150.00"),
            )

        inf = CatalogoInfraccion.objects.first()
        if inf:
            Infraccion.objects.create(
                acta=self.acta,
                tipo_infraccion=inf,
                descripcion="Falta de señalización visible",
                es_reincidencia=False,
            )

        self.client = APIClient()

    def _crear_usuario(self, rol, subtipo, prefijo):
        return crear_usuario(
            nombre_completo=f"{prefijo} {self.suffix}",
            email=f"{prefijo}.{self.suffix}@uubr.org.bo",
            password="claveSegura123",
            rol=Rol.objects.get(nombre=rol),
            subtipo=subtipo,
        )

    def test_generar_reporte_ejecutivo_pdf(self):
        self.client.force_authenticate(user=self.admin)
        payload = {
            "tipo": "ejecutivo",
            "formato": "PDF",
            "campania_id": self.campania.id,
        }
        res = self.client.post("/api/reportes/generar/", payload, format="json")
        self.assertEqual(res.status_code, status.HTTP_200_OK)
        self.assertEqual(res["Content-Type"], "application/pdf")
        self.assertIn("attachment; filename=", res["Content-Disposition"])
        self.assertTrue(res.content.startswith(b"%PDF-"))

        # Verifica persistencia del registro en BD
        reporte_id = int(res["X-Reporte-Id"])
        reporte = ReporteGenerado.objects.get(id=reporte_id)
        self.assertEqual(reporte.tipo, "ejecutivo")
        self.assertEqual(reporte.formato, "PDF")
        self.assertEqual(reporte.generado_por, self.admin)
        self.assertEqual(reporte.parametros.get("campania_id"), self.campania.id)

    def test_generar_reporte_ejecutivo_excel(self):
        self.client.force_authenticate(user=self.oficial)
        payload = {
            "tipo": "ejecutivo",
            "formato": "Excel",
            "campania_id": self.campania.id,
        }
        res = self.client.post("/api/reportes/generar/", payload, format="json")
        self.assertEqual(res.status_code, status.HTTP_200_OK)
        self.assertEqual(
            res["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        self.assertIn(".xlsx", res["Content-Disposition"])

        # Valida que sea un Excel legible con las dos hojas
        wb = openpyxl.load_workbook(io.BytesIO(res.content))
        self.assertIn("Resumen Ejecutivo", wb.sheetnames)
        self.assertIn("Inspecciones", wb.sheetnames)

    def test_generar_dictamen_tecnico_pdf(self):
        self.client.force_authenticate(user=self.admin)
        payload = {
            "tipo": "dictamen",
            "formato": "PDF",
            "acta_id": self.acta.id,
        }
        res = self.client.post("/api/reportes/generar/", payload, format="json")
        self.assertEqual(res.status_code, status.HTTP_200_OK)
        self.assertEqual(res["Content-Type"], "application/pdf")
        self.assertTrue(res.content.startswith(b"%PDF-"))

    def test_generar_dictamen_tecnico_excel(self):
        self.client.force_authenticate(user=self.admin)
        payload = {
            "tipo": "dictamen",
            "formato": "Excel",
            "acta_id": self.acta.id,
        }
        res = self.client.post("/api/reportes/generar/", payload, format="json")
        self.assertEqual(res.status_code, status.HTTP_200_OK)
        wb = openpyxl.load_workbook(io.BytesIO(res.content))
        self.assertIn("Dictamen Técnico", wb.sheetnames)

    def test_generar_dictamen_por_predio_id(self):
        self.client.force_authenticate(user=self.oficial)
        payload = {
            "tipo": "dictamen",
            "formato": "PDF",
            "predio_id": self.predio.id,
        }
        res = self.client.post("/api/reportes/generar/", payload, format="json")
        self.assertEqual(res.status_code, status.HTTP_200_OK)
        self.assertTrue(res.content.startswith(b"%PDF-"))

    def test_inspector_sin_permiso_web(self):
        self.client.force_authenticate(user=self.inspector)
        payload = {
            "tipo": "ejecutivo",
            "formato": "PDF",
            "campania_id": self.campania.id,
        }
        res = self.client.post("/api/reportes/generar/", payload, format="json")
        self.assertEqual(res.status_code, status.HTTP_403_FORBIDDEN)

    def test_validacion_parametros_faltantes(self):
        self.client.force_authenticate(user=self.admin)
        # Ejecutivo sin campania_id
        res1 = self.client.post(
            "/api/reportes/generar/",
            {"tipo": "ejecutivo", "formato": "PDF"},
            format="json",
        )
        self.assertEqual(res1.status_code, status.HTTP_400_BAD_REQUEST)

        # Dictamen sin acta_id ni predio_id
        res2 = self.client.post(
            "/api/reportes/generar/",
            {"tipo": "dictamen", "formato": "PDF"},
            format="json",
        )
        self.assertEqual(res2.status_code, status.HTTP_400_BAD_REQUEST)

    def test_recurso_no_encontrado(self):
        self.client.force_authenticate(user=self.admin)
        res = self.client.post(
            "/api/reportes/generar/",
            {"tipo": "ejecutivo", "formato": "PDF", "campania_id": 999999},
            format="json",
        )
        self.assertEqual(res.status_code, status.HTTP_404_NOT_FOUND)

    def test_historial_reportes_list(self):
        self.client.force_authenticate(user=self.admin)
        # Generar un reporte primero
        self.client.post(
            "/api/reportes/generar/",
            {"tipo": "ejecutivo", "formato": "PDF", "campania_id": self.campania.id},
            format="json",
        )

        res = self.client.get("/api/reportes/reportes/")
        self.assertEqual(res.status_code, status.HTTP_200_OK)
        self.assertGreaterEqual(len(res.data), 1)
        primer = res.data[0]
        self.assertIn("tipo", primer)
        self.assertIn("formato", primer)
        self.assertIn("generado_por_nombre", primer)
