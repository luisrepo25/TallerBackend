import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# Paleta institucional UUBR
COLOR_MARINO = "14213D"
COLOR_MARINO_SEC = "1F3357"
COLOR_BLANCO = "FFFFFF"
COLOR_GRIS_FONDO = "F4F5F7"
COLOR_BORDE = "D9D9D9"
COLOR_CRITICO = "FDECEE"
COLOR_BAJO = "E8F5E9"
COLOR_MEDIO = "FFF9C4"

FONT_TITULO = Font(name="Calibri", size=16, bold=True, color=COLOR_BLANCO)
FONT_SUBTITULO = Font(name="Calibri", size=12, bold=True, color=COLOR_BLANCO)
FONT_SECCION = Font(name="Calibri", size=11, bold=True, color=COLOR_MARINO)
FONT_CABECERA = Font(name="Calibri", size=10, bold=True, color=COLOR_BLANCO)
FONT_NORMAL = Font(name="Calibri", size=10, color="000000")
FONT_BOLD = Font(name="Calibri", size=10, bold=True, color="000000")

FILL_HEADER = PatternFill(start_color=COLOR_MARINO, end_color=COLOR_MARINO, fill_type="solid")
FILL_SUBHEADER = PatternFill(
    start_color=COLOR_MARINO_SEC, end_color=COLOR_MARINO_SEC, fill_type="solid"
)
FILL_ZEBRA = PatternFill(
    start_color=COLOR_GRIS_FONDO, end_color=COLOR_GRIS_FONDO, fill_type="solid"
)

BORDER_THIN = Border(
    left=Side(style="thin", color=COLOR_BORDE),
    right=Side(style="thin", color=COLOR_BORDE),
    top=Side(style="thin", color=COLOR_BORDE),
    bottom=Side(style="thin", color=COLOR_BORDE),
)

ALIGN_CENTER = Alignment(horizontal="center", vertical="center")
ALIGN_LEFT = Alignment(horizontal="left", vertical="center")
ALIGN_RIGHT = Alignment(horizontal="right", vertical="center")


def _autoajustar_columnas(ws):
    """Ajusta automaticamente el ancho de las columnas con un margen minimo."""
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val = str(cell.value or "")
            if "\n" in val:
                val = max(val.split("\n"), key=len)
            max_len = max(max_len, len(val))
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)


def generar_excel_ejecutivo(datos: dict) -> bytes:
    """Genera un archivo Excel (.xlsx) con el reporte ejecutivo de una campaña.

    Incluye Hoja 1 (Resumen ejecutivo y métricas de riesgo) y Hoja 2 (Detalle de predios e
    inspecciones).
    """
    wb = Workbook()

    # ----------------------------------------------------
    # HOJA 1: RESUMEN EJECUTIVO
    # ----------------------------------------------------
    ws_resumen = wb.active
    ws_resumen.title = "Resumen Ejecutivo"
    ws_resumen.views.sheetView[0].showGridLines = True

    # Titulo Principal
    ws_resumen.merge_cells("A1:F1")
    c1 = ws_resumen["A1"]
    c1.value = "CUERPO DE BOMBEROS VOLUNTARIOS UUBR - SANTA CRUZ"
    c1.font = FONT_TITULO
    c1.fill = FILL_HEADER
    c1.alignment = ALIGN_CENTER
    ws_resumen.row_dimensions[1].height = 30

    ws_resumen.merge_cells("A2:F2")
    c2 = ws_resumen["A2"]
    c2.value = f"REPORTE EJECUTIVO DE CAMPAÑA: {datos['campania']['nombre']}"
    c2.font = FONT_SUBTITULO
    c2.fill = FILL_SUBHEADER
    c2.alignment = ALIGN_CENTER
    ws_resumen.row_dimensions[2].height = 24

    fila = 4
    # Metadatos de la Campaña
    ws_resumen.cell(row=fila, column=1, value="1. INFORMACIÓN GENERAL DE LA CAMPAÑA").font = (
        FONT_SECCION
    )
    fila += 1

    meta_items = [
        ("Campaña:", datos["campania"]["nombre"]),
        ("Estado:", datos["campania"]["estado"].capitalize()),
        (
            "Fechas de vigencia:",
            f"{datos['campania']['fecha_inicio']} al {datos['campania']['fecha_fin']}",
        ),
        ("Creado por:", datos["campania"]["creado_por"]),
        (
            "Inspectores asignados:",
            ", ".join(datos["campania"]["inspectores"]) or "Sin inspectores asignados",
        ),
    ]
    for label, val in meta_items:
        ws_resumen.cell(row=fila, column=1, value=label).font = FONT_BOLD
        ws_resumen.cell(row=fila, column=2, value=val).font = FONT_NORMAL
        fila += 1

    fila += 1
    # Métricas de Avance Territorial
    ws_resumen.cell(row=fila, column=1, value="2. AVANCE TERRITORIAL").font = FONT_SECCION
    fila += 1

    metricas_avance = [
        ("Predios dentro del polígono:", datos["avance"]["total_predios"]),
        ("Predios con inspección:", datos["avance"]["total_inspeccionados"]),
        ("Actas cerradas (completadas):", datos["avance"]["actas_cerradas"]),
        ("Actas en curso / abiertas:", datos["avance"]["actas_abiertas"]),
        ("Porcentaje de avance:", f"{datos['avance']['porcentaje_avance']}%"),
    ]
    for label, val in metricas_avance:
        ws_resumen.cell(row=fila, column=1, value=label).font = FONT_BOLD
        ws_resumen.cell(row=fila, column=2, value=val).font = FONT_NORMAL
        fila += 1

    fila += 1
    # Distribución de Riesgo (K-Means / Scoring)
    ws_resumen.cell(row=fila, column=1, value="3. EVALUACIÓN Y DISTRIBUCIÓN DE RIESGO").font = (
        FONT_SECCION
    )
    fila += 1

    headers_riesgo = ["Nivel de Riesgo (Cluster)", "Predios", "Porcentaje (%)"]
    for col_idx, h in enumerate(headers_riesgo, start=1):
        cell = ws_resumen.cell(row=fila, column=col_idx, value=h)
        cell.font = FONT_CABECERA
        cell.fill = FILL_HEADER
        cell.alignment = ALIGN_CENTER
        cell.border = BORDER_THIN
    fila += 1

    clusters_info = [
        ("Crítico", datos["riesgo"]["critico"], datos["riesgo"]["pct_critico"]),
        ("Medio", datos["riesgo"]["medio"], datos["riesgo"]["pct_medio"]),
        ("Bajo", datos["riesgo"]["bajo"], datos["riesgo"]["pct_bajo"]),
        (
            "Sin clasificar",
            datos["riesgo"]["sin_clasificar"],
            datos["riesgo"]["pct_sin_clasificar"],
        ),
    ]
    for nombre_cluster, cantidad, pct in clusters_info:
        c_nom = ws_resumen.cell(row=fila, column=1, value=nombre_cluster)
        c_cnt = ws_resumen.cell(row=fila, column=2, value=cantidad)
        c_pct = ws_resumen.cell(row=fila, column=3, value=f"{pct}%")

        c_nom.font = FONT_BOLD
        c_cnt.font = FONT_NORMAL
        c_pct.font = FONT_NORMAL

        c_cnt.alignment = ALIGN_RIGHT
        c_pct.alignment = ALIGN_RIGHT

        c_nom.border = BORDER_THIN
        c_cnt.border = BORDER_THIN
        c_pct.border = BORDER_THIN

        if nombre_cluster == "Crítico" and cantidad > 0:
            c_nom.fill = PatternFill(
                start_color=COLOR_CRITICO, end_color=COLOR_CRITICO, fill_type="solid"
            )
        fila += 1

    fila += 1
    # Métricas Globales
    ws_resumen.cell(row=fila, column=1, value="Scoring de riesgo promedio:").font = FONT_BOLD
    ws_resumen.cell(
        row=fila, column=2, value=f"{datos['riesgo']['scoring_promedio']} / 100"
    ).font = FONT_NORMAL
    fila += 1

    ws_resumen.cell(row=fila, column=1, value="Carga de fuego promedio:").font = FONT_BOLD
    ws_resumen.cell(
        row=fila, column=2, value=f"{datos['riesgo']['carga_fuego_promedio']} MJ/m²"
    ).font = FONT_NORMAL
    fila += 1

    ws_resumen.cell(row=fila, column=1, value="Distancia promedio a hidrante:").font = FONT_BOLD
    ws_resumen.cell(
        row=fila, column=2, value=f"{datos['riesgo']['distancia_hidrante_promedio']} m"
    ).font = FONT_NORMAL
    fila += 2

    # Infracciones más comunes
    if datos["infracciones"]:
        ws_resumen.cell(
            row=fila, column=1, value="4. INFRACCIONES MÁS FRECUENTES EN LA CAMPAÑA"
        ).font = FONT_SECCION
        fila += 1

        inf_headers = ["Tipo de Infracción", "Casos Detectados", "Reincidencias"]
        for col_idx, h in enumerate(inf_headers, start=1):
            cell = ws_resumen.cell(row=fila, column=col_idx, value=h)
            cell.font = FONT_CABECERA
            cell.fill = FILL_HEADER
            cell.alignment = ALIGN_CENTER
            cell.border = BORDER_THIN
        fila += 1

        for inf in datos["infracciones"][:10]:
            c_nom = ws_resumen.cell(row=fila, column=1, value=inf["nombre"])
            c_tot = ws_resumen.cell(row=fila, column=2, value=inf["total"])
            c_rei = ws_resumen.cell(row=fila, column=3, value=inf["reincidentes"])

            for c in (c_nom, c_tot, c_rei):
                c.font = FONT_NORMAL
                c.border = BORDER_THIN
            c_tot.alignment = ALIGN_RIGHT
            c_rei.alignment = ALIGN_RIGHT
            fila += 1

    _autoajustar_columnas(ws_resumen)

    # ----------------------------------------------------
    # HOJA 2: DETALLE TABULAR DE INSPECCIONES
    # ----------------------------------------------------
    ws_detalle = wb.create_sheet(title="Inspecciones")
    ws_detalle.views.sheetView[0].showGridLines = True

    columnas_detalle = [
        "Acta ID",
        "Predio",
        "Dirección",
        "Tipo Inmueble",
        "Inspector",
        "Fecha Inspección",
        "Carga Fuego (MJ/m²)",
        "Scoring (0-100)",
        "Cluster Riesgo",
        "Extintores",
        "Rutas Obstruidas",
        "Fallas Eléctricas",
        "Infracciones",
        "Estado Acta",
    ]

    for col_idx, col_name in enumerate(columnas_detalle, start=1):
        cell = ws_detalle.cell(row=1, column=col_idx, value=col_name)
        cell.font = FONT_CABECERA
        cell.fill = FILL_HEADER
        cell.alignment = ALIGN_CENTER
        cell.border = BORDER_THIN
    ws_detalle.row_dimensions[1].height = 24

    for r_idx, reg in enumerate(datos["detalles_actas"], start=2):
        row_cells = [
            ws_detalle.cell(row=r_idx, column=1, value=reg["acta_id"]),
            ws_detalle.cell(row=r_idx, column=2, value=reg["predio_nombre"]),
            ws_detalle.cell(row=r_idx, column=3, value=reg["direccion"]),
            ws_detalle.cell(row=r_idx, column=4, value=reg["tipo_predio"]),
            ws_detalle.cell(row=r_idx, column=5, value=reg["inspector_nombre"]),
            ws_detalle.cell(row=r_idx, column=6, value=reg["fecha_inspeccion"]),
            ws_detalle.cell(row=r_idx, column=7, value=reg["carga_fuego"]),
            ws_detalle.cell(row=r_idx, column=8, value=reg["scoring_riesgo"]),
            ws_detalle.cell(row=r_idx, column=9, value=reg["cluster_riesgo"]),
            ws_detalle.cell(row=r_idx, column=10, value=reg["estado_extintores"]),
            ws_detalle.cell(
                row=r_idx,
                column=11,
                value="Sí" if reg["rutas_obstruidas"] else "No",
            ),
            ws_detalle.cell(
                row=r_idx,
                column=12,
                value="Sí" if reg["fallas_electricas"] else "No",
            ),
            ws_detalle.cell(row=r_idx, column=13, value=reg["infracciones_conteo"]),
            ws_detalle.cell(
                row=r_idx,
                column=14,
                value="Cerrada" if reg["cerrada"] else "En curso",
            ),
        ]

        zebra = FILL_ZEBRA if r_idx % 2 == 0 else None
        for cell in row_cells:
            cell.font = FONT_NORMAL
            cell.border = BORDER_THIN
            if zebra:
                cell.fill = zebra

        # Alinear columnas numericas
        row_cells[0].alignment = ALIGN_CENTER
        row_cells[5].alignment = ALIGN_CENTER
        row_cells[6].alignment = ALIGN_RIGHT
        row_cells[7].alignment = ALIGN_RIGHT
        row_cells[8].alignment = ALIGN_CENTER
        row_cells[10].alignment = ALIGN_CENTER
        row_cells[11].alignment = ALIGN_CENTER
        row_cells[12].alignment = ALIGN_RIGHT
        row_cells[13].alignment = ALIGN_CENTER

    _autoajustar_columnas(ws_detalle)

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def generar_excel_dictamen(datos: dict) -> bytes:
    """Genera un archivo Excel (.xlsx) con el dictamen técnico de inspección de un
    predio."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Dictamen Técnico"
    ws.views.sheetView[0].showGridLines = True

    # Titulo Principal
    ws.merge_cells("A1:E1")
    c1 = ws["A1"]
    c1.value = "CUERPO DE BOMBEROS VOLUNTARIOS UUBR - SANTA CRUZ"
    c1.font = FONT_TITULO
    c1.fill = FILL_HEADER
    c1.alignment = ALIGN_CENTER
    ws.row_dimensions[1].height = 28

    ws.merge_cells("A2:E2")
    c2 = ws["A2"]
    c2.value = f"DICTAMEN TÉCNICO DE INSPECCIÓN · ACTA #{datos['acta']['id']}"
    c2.font = FONT_SUBTITULO
    c2.fill = FILL_SUBHEADER
    c2.alignment = ALIGN_CENTER
    ws.row_dimensions[2].height = 22

    fila = 4
    # 1. Datos del Inmueble
    ws.cell(row=fila, column=1, value="1. DATOS DEL INMUEBLE / PREDIO").font = FONT_SECCION
    fila += 1

    datos_predio = [
        ("Nombre del predio:", datos["predio"]["nombre"]),
        ("Dirección:", datos["predio"]["direccion"]),
        ("Tipo de edificación:", datos["predio"]["tipo"]),
        ("Área construida:", f"{datos['predio']['area_m2']} m²"),
        ("Aforo máximo:", str(datos["predio"]["aforo"] or "No especificado")),
        (
            "Hidrante más cercano:",
            f"{datos['predio']['hidrante_codigo']} ({datos['predio']['distancia_hidrante_m']} m)",
        ),
    ]
    for label, val in datos_predio:
        ws.cell(row=fila, column=1, value=label).font = FONT_BOLD
        ws.cell(row=fila, column=2, value=val).font = FONT_NORMAL
        fila += 1

    fila += 1
    # 2. Datos de la Inspección
    ws.cell(row=fila, column=1, value="2. REGISTRO DE INSPECCIÓN EN CAMPO").font = FONT_SECCION
    fila += 1

    datos_insp = [
        ("Campaña asociada:", datos["acta"]["campania_nombre"]),
        ("Inspector actuante:", datos["acta"]["inspector_nombre"]),
        ("Fecha y hora de inspección:", datos["acta"]["fecha_inspeccion"]),
        (
            "Estado del acta:",
            "Cerrada (Definitiva)" if datos["acta"]["cerrada"] else "En curso (Preliminar)",
        ),
    ]
    for label, val in datos_insp:
        ws.cell(row=fila, column=1, value=label).font = FONT_BOLD
        ws.cell(row=fila, column=2, value=val).font = FONT_NORMAL
        fila += 1

    fila += 1
    # 3. Materiales Combustibles y Carga de Fuego NB 58005
    ws.cell(
        row=fila,
        column=1,
        value="3. CARGA DE FUEGO (NORMA BOLIVIANA NB 58005)",
    ).font = FONT_SECCION
    fila += 1

    mat_headers = [
        "Material Combustible",
        "Peso (kg)",
        "Hi (MJ/kg)",
        "Peligrosidad (Ci)",
        "Calor Aportado (MJ)",
    ]
    for col_idx, h in enumerate(mat_headers, start=1):
        cell = ws.cell(row=fila, column=col_idx, value=h)
        cell.font = FONT_CABECERA
        cell.fill = FILL_HEADER
        cell.alignment = ALIGN_CENTER
        cell.border = BORDER_THIN
    fila += 1

    for mat in datos["materiales"]:
        c_nom = ws.cell(row=fila, column=1, value=mat["nombre"])
        c_pes = ws.cell(row=fila, column=2, value=mat["peso_kg"])
        c_hi = ws.cell(row=fila, column=3, value=mat["poder_calorifico"])
        c_ci = ws.cell(row=fila, column=4, value=mat["coeficiente_peligrosidad"])
        c_tot = ws.cell(row=fila, column=5, value=mat["calor_total_mj"])

        for c in (c_nom, c_pes, c_hi, c_ci, c_tot):
            c.font = FONT_NORMAL
            c.border = BORDER_THIN
        c_pes.alignment = ALIGN_RIGHT
        c_hi.alignment = ALIGN_RIGHT
        c_ci.alignment = ALIGN_RIGHT
        c_tot.alignment = ALIGN_RIGHT
        fila += 1

    ws.cell(row=fila, column=1, value="Carga de fuego resultante:").font = FONT_BOLD
    ws.cell(row=fila, column=2, value=f"{datos['acta']['carga_fuego']} MJ/m²").font = FONT_BOLD
    fila += 2

    # 4. Infracciones y Medidas de Seguridad
    ws.cell(row=fila, column=1, value="4. CONDICIONES DE SEGURIDAD E INFRACCIONES").font = (
        FONT_SECCION
    )
    fila += 1

    seguridad_items = [
        ("Estado de extintores:", datos["acta"]["estado_extintores"]),
        (
            "Rutas de evacuación obstruidas:",
            "Sí (Grave)" if datos["acta"]["rutas_obstruidas"] else "No (Despejadas)",
        ),
        (
            "Fallas eléctricas detectadas:",
            "Sí (Riesgo alto)" if datos["acta"]["fallas_electricas"] else "No (En orden)",
        ),
    ]
    for label, val in seguridad_items:
        ws.cell(row=fila, column=1, value=label).font = FONT_BOLD
        ws.cell(row=fila, column=2, value=val).font = FONT_NORMAL
        fila += 1

    if datos["infracciones"]:
        fila += 1
        inf_headers = ["Infracción Registrada", "Severidad", "Reincidencia", "Observaciones"]
        for col_idx, h in enumerate(inf_headers, start=1):
            cell = ws.cell(row=fila, column=col_idx, value=h)
            cell.font = FONT_CABECERA
            cell.fill = FILL_HEADER
            cell.alignment = ALIGN_CENTER
            cell.border = BORDER_THIN
        fila += 1

        for inf in datos["infracciones"]:
            c_nom = ws.cell(row=fila, column=1, value=inf["nombre"])
            c_sev = ws.cell(row=fila, column=2, value=inf["severidad"])
            c_rei = ws.cell(row=fila, column=3, value="REINCIDENTE" if inf["reincidente"] else "No")
            c_det = ws.cell(row=fila, column=4, value=inf["descripcion"])

            for c in (c_nom, c_sev, c_rei, c_det):
                c.font = FONT_NORMAL
                c.border = BORDER_THIN
            c_sev.alignment = ALIGN_CENTER
            c_rei.alignment = ALIGN_CENTER
            fila += 1

    fila += 2
    # 5. Dictamen y Conclusión
    ws.cell(row=fila, column=1, value="5. DICTAMEN FINAL DE RIESGO").font = FONT_SECCION
    fila += 1

    ws.cell(row=fila, column=1, value="Scoring de Riesgo (0 - 100):").font = FONT_BOLD
    ws.cell(row=fila, column=2, value=f"{datos['acta']['scoring_riesgo']} pts").font = FONT_BOLD
    fila += 1

    ws.cell(row=fila, column=1, value="Clasificación de Riesgo (Cluster):").font = FONT_BOLD
    c_riesgo = ws.cell(row=fila, column=2, value=datos["acta"]["cluster_riesgo"])
    c_riesgo.font = FONT_BOLD
    if datos["acta"]["cluster_riesgo"] == "Critico":
        c_riesgo.fill = PatternFill(
            start_color=COLOR_CRITICO, end_color=COLOR_CRITICO, fill_type="solid"
        )
    elif datos["acta"]["cluster_riesgo"] == "Bajo":
        c_riesgo.fill = PatternFill(start_color=COLOR_BAJO, end_color=COLOR_BAJO, fill_type="solid")
    elif datos["acta"]["cluster_riesgo"] == "Medio":
        c_riesgo.fill = PatternFill(
            start_color=COLOR_MEDIO, end_color=COLOR_MEDIO, fill_type="solid"
        )

    _autoajustar_columnas(ws)

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
